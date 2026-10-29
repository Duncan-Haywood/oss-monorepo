#!/usr/bin/env python3
"""Build Raphael's Map's U.S. data files from the raw official downloads.

    python3 -I pipeline/build_us.py --raw /path/to/raw --out ventures/raphaels-map/data

Reads the files pipeline/fetch.py put on the `raphaels-map-raw` branch
(raw/manifest.json lists them) and writes, exactly as ARCHITECTURE.md
describes:

    data/us/geographies.json           counties, states, CBSAs
    data/us/county_observations.json   every county x every county metric
    data/us/state_observations.json    same metrics, state rows of the ACS
    data/us/cbsa_observations.json     same metrics, CBSA rows of the ACS
    data/us/build_checks.json          internal consistency checks of this build
    data/geo/us_counties.json          national county + state SVG paths
    data/geo/state_<fips>.json         one file per state, DC and Puerto Rico
    data/sources.json                  source registry (+ pipeline/registry/*.json)
    data/metrics.json                  metric dictionary (+ pipeline/registry/*.json)

Rules this code enforces (SPEC section 10, "Data integrity"):

* ACS cells are looked up by their published label and the build stops if
  the codes the spec names do not carry the expected labels.
* A value that is not in a source file is null with a flag, never 0.
* FIPS codes are strings with leading zeros.
* The Religion Census join rejects duplicate and ambiguous county rows.
* Metros come from the ACS's own CBSA rows; they are never summed from
  county rows.

Only openpyxl (Excel) and pyshp (shapefile) are needed beyond the standard
library.
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import math
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
ACS_YEAR = 2024
ACS_PERIOD = "2020-2024"
ACS_VINTAGE_LABEL = "2020–2024 American Community Survey 5-year"
USRC_VINTAGE_LABEL = "2020 U.S. Religion Census"


class BuildError(RuntimeError):
    """A source file does not look the way the build was written against."""


# ---------------------------------------------------------------- observations

def obs(e, moe=None, flags=()):
    """The Observation object of ARCHITECTURE.md."""
    out = []
    for f in flags:
        if f not in out:
            out.append(f)
    return {"e": e, "moe": moe, "flags": out}


# ACS special values. In the table-based Summary File a negative number in an
# estimate or margin-of-error column is a code, not a value (ACS Summary File
# documentation, "Explanation of symbols"). The only code that carries a
# number is -555555555 in a margin-of-error column: the estimate is
# controlled to the official population estimates, so a sampling margin of
# error is not appropriate and the Census Bureau's data tables show it as 0
# ("*****"). Every other code makes the value null.
SENTINEL_MOE_CONTROLLED = -555555555
SENTINEL_ESTIMATE_FLAGS = {
    -666666666: "suppressed_small_sample",  # too few sample observations to compute an estimate
    -999999999: "not_published",            # estimate not displayed / not available
    -888888888: "not_published",            # not applicable
    -222222222: "not_published",
    -333333333: "not_published",
    -555555555: "not_published",
}


def parse_cell(est_raw: str, moe_raw: str | None):
    """(estimate, moe, flags) from one Summary File estimate/MOE pair.

    Never turns a special code into a number; returns ints for real values."""
    flags: list[str] = []
    try:
        e = int(est_raw)
    except (TypeError, ValueError):
        e = None
        flags.append("not_published")
    if e is not None and e < 0:
        flags.append(SENTINEL_ESTIMATE_FLAGS.get(e, "not_published"))
        e = None
    moe = None
    if moe_raw not in (None, ""):
        try:
            m = int(moe_raw)
        except ValueError:
            m = None
        if m is not None:
            if m == SENTINEL_MOE_CONTROLLED:
                moe = 0
            elif m >= 0:
                moe = m
            # other negative codes (-222222222 too few observations,
            # -333333333 open-ended median, -999999999) -> moe stays None
    if e is None:
        moe = None
    return e, moe, flags


def sum_obs(parts: list[dict]) -> dict:
    """Sum of published cells. MOE is the Census Bureau's approximation
    sqrt(sum(moe_i^2)), flagged moe_approx; when several parts are zero
    estimates only the largest of their MOEs is used (ACS handbook guidance
    for sums that include multiple zero estimates)."""
    if any(p["e"] is None for p in parts):
        fl = [f for p in parts for f in p["flags"] if f != "moe_approx"] or ["not_published"]
        return obs(None, None, fl)
    e = sum(p["e"] for p in parts)
    if any(p["moe"] is None for p in parts):
        moe = None
    else:
        nonzero = [p["moe"] for p in parts if p["e"] != 0]
        zero = [p["moe"] for p in parts if p["e"] == 0]
        terms = nonzero + ([max(zero)] if zero else [])
        moe = round(math.sqrt(sum(m * m for m in terms)))
    flags = [f for p in parts for f in p["flags"]]
    if len(parts) > 1:
        flags.append("moe_approx")
    return obs(e, moe, flags)


def ratio_obs(num: dict, den: dict, scale: float, *, subset: bool, digits: int = 1,
              extra_flags=()) -> dict:
    """num/den*scale, flagged `derived`. MOE by the Census Bureau's formulas
    for derived proportions (subset=True: numerator is a subset of the
    denominator) or ratios, flagged moe_approx; null when an input is null."""
    flags = ["derived", *extra_flags]
    if num["e"] is None or den["e"] is None:
        why = [f for f in num["flags"] + den["flags"] if f not in ("moe_approx", "derived")]
        return obs(None, None, flags + (why or ["not_published"]))
    if den["e"] == 0:
        return obs(None, None, flags + ["not_published"])
    r = num["e"] / den["e"]
    moe = None
    if num.get("moe") is not None and den.get("moe") is not None:
        mx, my, y = num["moe"], den["moe"], den["e"]
        under = mx * mx - r * r * my * my if subset else -1
        if under < 0:
            under = mx * mx + r * r * my * my
        moe = round(math.sqrt(under) / y * scale, digits)
        flags.append("moe_approx")
    return obs(round(r * scale, digits), moe, flags)


# ------------------------------------------------------------------ ACS cells

# The codes SPEC 5.1 names, with the label each must carry in the published
# 2024 group metadata. A mismatch stops the build.
SPEC_CODES = {
    "B01003_001E": "Estimate!!Total",
    "B01001_026E": "Estimate!!Total:!!Female:",
    "B01001_031E": "Estimate!!Total:!!Female:!!18 and 19 years",
    "B01001_032E": "Estimate!!Total:!!Female:!!20 years",
    "B01001_033E": "Estimate!!Total:!!Female:!!21 years",
    "B01001_034E": "Estimate!!Total:!!Female:!!22 to 24 years",
    "B01001_035E": "Estimate!!Total:!!Female:!!25 to 29 years",
    "B01001_036E": "Estimate!!Total:!!Female:!!30 to 34 years",
}

# Never-married women by age come from B12002, found by label:
# Female -> Never married -> age. Only bins the table publishes are used.
B12002_NEVER_MARRIED_AGES = {
    "20_24": "20 to 24 years",
    "25_29": "25 to 29 years",
    "30_34": "30 to 34 years",
}
B12002_PREFIX = "Estimate!!Total:!!Female:!!Never married:!!"

ACS_SIMPLE = {  # metric -> list of codes summed
    "total_population": ["B01003_001E"],
    "female_population": ["B01001_026E"],
    "women_18_19": ["B01001_031E"],
    "women_20_24": ["B01001_032E", "B01001_033E", "B01001_034E"],
    "women_25_29": ["B01001_035E"],
    "women_30_34": ["B01001_036E"],
    "women_20_29": ["B01001_032E", "B01001_033E", "B01001_034E", "B01001_035E"],
}


def load_group_labels(path: Path) -> dict[str, str]:
    d = json.loads(Path(path).read_text(encoding="utf-8"))
    v = d.get("variables", d)
    return {k: (val.get("label") or "") for k, val in v.items() if isinstance(val, dict)}


def find_by_label(labels: dict[str, str], label: str) -> str:
    hits = [k for k, v in labels.items() if v == label and k.endswith("E")]
    if len(hits) != 1:
        raise BuildError(f"ACS label {label!r}: expected exactly one estimate cell, found {hits}")
    return hits[0]


def load_shell_labels(path: Path) -> dict[str, str]:
    """Table-shell file -> {'B01001_032': 'Total:!!Female:!!20 years'} built
    from the indent levels, so it can be compared with the API labels."""
    out: dict[str, str] = {}
    stack: dict[str, list[str]] = defaultdict(list)
    with open(path, encoding="utf-8", errors="replace") as fh:
        header = fh.readline().rstrip("\n").split("|")
        ix = {h: i for i, h in enumerate(header)}
        for line in fh:
            f = line.rstrip("\n").split("|")
            if len(f) < len(header):
                continue
            tid, uid, label = f[ix["Table ID"]], f[ix["Unique ID"]], f[ix["Label"]]
            try:
                indent = int(float(f[ix["Indent"]]))
            except ValueError:
                continue
            st = stack[tid]
            del st[indent:]
            st.append(label.strip())
            out[uid] = "!!".join(st)
    return out


def resolve_acs_cells(meta_dir: Path, shells: dict[str, str] | None = None) -> dict:
    """Check every code against its published label; return the metric ->
    codes map (including the B12002 cells found by label)."""
    labels = {}
    for g in ("B01001", "B12002", "B01003"):
        p = Path(meta_dir) / f"groups_{g}.json"
        if not p.exists():
            raise BuildError(f"missing ACS group metadata {p}")
        labels.update(load_group_labels(p))
    for code, want in SPEC_CODES.items():
        got = labels.get(code)
        if got != want:
            raise BuildError(f"ACS code {code} has label {got!r}, the spec expects {want!r}; "
                             "the build will not guess which cell is meant")
    cells = {k: list(v) for k, v in ACS_SIMPLE.items()}
    for key, age in B12002_NEVER_MARRIED_AGES.items():
        cells[f"never_married_women_{key}"] = [find_by_label(labels, B12002_PREFIX + age)]
    cells["never_married_women_20_29"] = cells["never_married_women_20_24"] + cells["never_married_women_25_29"]
    # Check-only cells: every B12002 female cell for ages 20-29, any marital
    # status (the leaves under Female). Their sum must equal B01001 women 20-29
    # for the never-married share to have one consistent denominator.
    check = sorted(k for k, v in labels.items() if k.startswith("B12002_") and k.endswith("E")
                   and v.startswith("Estimate!!Total:!!Female:!!")
                   and (v.endswith("!!20 to 24 years") or v.endswith("!!25 to 29 years")))
    if len(check) != 12:
        raise BuildError(f"B12002: expected 12 female 20-29 cells (6 marital statuses x 2 ages), found {len(check)}")
    cells["_check_b12002_female_20_29"] = check
    if shells is not None:  # second, independent check: the Summary File table shells
        for codes in cells.values():
            for code in codes:
                uid = code[:-1]
                api = labels[code].removeprefix("Estimate!!")
                if shells.get(uid) is None:
                    raise BuildError(f"{uid} missing from the table shells")
                if shells[uid].rstrip(":") != api.rstrip(":"):
                    raise BuildError(f"{uid}: shell label {shells[uid]!r} != API label {api!r}")
    return {"cells": cells, "labels": {c: labels[c] for m, cs in cells.items() if not m.startswith("_") for c in cs}}


def sf_columns(code: str) -> tuple[str, str]:
    """'B01001_032E' -> ('B01001_E032', 'B01001_M032') (table-based SF names)."""
    m = re.fullmatch(r"([A-Z]\d{5}[A-Z]*)_(\d{3})E", code)
    if not m:
        raise BuildError(f"not an ACS estimate code: {code}")
    return f"{m.group(1)}_E{m.group(2)}", f"{m.group(1)}_M{m.group(2)}"


def read_sf(path: Path, codes: list[str]) -> dict[str, dict[str, tuple]]:
    """Read a table-based Summary File .dat: {GEO_ID: {code: (e, moe, flags)}}."""
    rows: dict[str, dict[str, tuple]] = {}
    with open(path, encoding="utf-8") as fh:
        header = fh.readline().rstrip("\r\n").split("|")
        ix = {h: i for i, h in enumerate(header)}
        if "GEO_ID" not in ix:
            raise BuildError(f"{path}: no GEO_ID column")
        cols = {}
        for c in codes:
            ce, cm = sf_columns(c)
            if ce not in ix:
                raise BuildError(f"{path}: column {ce} for {c} not found")
            cols[c] = (ix[ce], ix.get(cm))
        for line in fh:
            f = line.rstrip("\r\n").split("|")
            if len(f) != len(header):
                raise BuildError(f"{path}: ragged row {f[:1]}")
            g = f[ix["GEO_ID"]]
            if g in rows:
                raise BuildError(f"{path}: duplicate GEO_ID {g}")
            rows[g] = {c: parse_cell(f[ie], f[im] if im is not None else None) for c, (ie, im) in cols.items()}
    return rows


def geo_kind(geo_id: str):
    """GEO_ID -> (level, key) for the levels used; None for others."""
    if geo_id == "0100000US":
        return ("us", "US")
    m = re.fullmatch(r"0400000US(\d{2})", geo_id)
    if m:
        return ("state", m.group(1))
    m = re.fullmatch(r"0500000US(\d{5})", geo_id)
    if m:
        return ("county", m.group(1))
    m = re.fullmatch(r"310M\d00US(\d{5})", geo_id)
    if m:
        return ("cbsa", m.group(1))
    return None


def acs_observations(sf_dir: Path, meta_dir: Path) -> tuple[dict, dict]:
    """{level: {key: {metric: Observation}}} from the Summary File rows of
    each level (metros from CBSA rows, never from counties)."""
    shells_p = Path(sf_dir) / f"ACS{ACS_YEAR}5YR_Table_Shells.txt"
    shells = load_shell_labels(shells_p) if shells_p.exists() else None
    resolved = resolve_acs_cells(meta_dir, shells)
    cells = resolved["cells"]
    by_table: dict[str, set] = defaultdict(set)
    for codes in cells.values():
        for c in codes:
            by_table[c.split("_")[0]].add(c)
    # B12002 leaf cells for female 20-24 / 25-29 across marital statuses: used
    # only to check that B12002 and B01001 describe the same women.
    raw: dict[str, dict] = defaultdict(dict)
    for table, codes in by_table.items():
        p = Path(sf_dir) / f"acsdt5y{ACS_YEAR}-{table.lower()}.dat"
        if not p.exists():
            raise BuildError(f"missing Summary File {p}")
        for g, vals in read_sf(p, sorted(codes)).items():
            raw[g].update(vals)
    out: dict[str, dict] = {"us": {}, "state": {}, "county": {}, "cbsa": {}}
    mismatch: list[str] = []
    for g, vals in raw.items():
        kind = geo_kind(g)
        if not kind:
            continue
        level, key = kind
        o = {}
        for metric, codes in cells.items():
            if metric.startswith("_"):
                continue
            parts = []
            for c in codes:
                if c not in vals:
                    parts = None
                    break
                e, moe, fl = vals[c]
                parts.append(obs(e, moe, fl))
            if parts is None:
                o[metric] = obs(None, None, ["not_published"])
            elif len(parts) == 1:
                o[metric] = parts[0]
            else:
                o[metric] = sum_obs(parts)
        o["never_married_share_women_20_29"] = ratio_obs(
            o["never_married_women_20_29"], o["women_20_29"], 100, subset=True)
        o["women_20_29_per_1000"] = ratio_obs(o["women_20_29"], o["total_population"], 1000, subset=True)
        chk = [vals[c][0] for c in cells["_check_b12002_female_20_29"] if c in vals]
        if o["women_20_29"]["e"] is not None and len(chk) == 12 and None not in chk and sum(chk) != o["women_20_29"]["e"]:
            mismatch.append(g)
            o["never_married_share_women_20_29"]["flags"].append("unvalidated")
        out[level][key] = o
    resolved["b01001_b12002_female_20_29_mismatches"] = mismatch
    return out, resolved


# ---------------------------------------------------------- geography tables

def read_geos(path: Path) -> dict[str, dict]:
    """Geos20245YR.txt -> {GEO_ID: {'name':..., 'stusab':...}}."""
    out = {}
    with open(path, encoding="utf-8") as fh:
        header = fh.readline().rstrip("\r\n").split("|")
        ix = {h: i for i, h in enumerate(header)}
        for line in fh:
            f = line.rstrip("\r\n").split("|")
            if len(f) < len(header):
                continue
            out[f[ix["GEO_ID"]]] = {"name": f[ix["NAME"]], "stusab": f[ix["STUSAB"]]}
    return out


def read_gazetteer(path: Path) -> dict[str, dict]:
    z = zipfile.ZipFile(path)
    name = next(n for n in z.namelist() if n.endswith(".txt"))
    text = z.read(name).decode("utf-8")
    lines = text.splitlines()
    header = [h.strip() for h in lines[0].split("\t")]
    ix = {h: i for i, h in enumerate(header)}
    out = {}
    for line in lines[1:]:
        f = [x.strip() for x in line.split("\t")]
        if len(f) < len(header):
            continue
        geoid = f[ix["GEOID"]]
        if not re.fullmatch(r"\d{5}", geoid):
            raise BuildError(f"gazetteer GEOID {geoid!r} is not 5 digits")
        out[geoid] = {"lat": float(f[ix["INTPTLAT"]]), "lon": float(f[ix["INTPTLONG"]]),
                      "land_sqmi": float(f[ix["ALAND_SQMI"]])}
    return out


DELINEATION_HEADERS = ["CBSA Code", "CBSA Title", "Metropolitan/Micropolitan Statistical Area",
                       "FIPS State Code", "FIPS County Code"]


def read_delineation(path: Path) -> dict[str, dict]:
    """OMB/Census list1 delineation -> {county_fips: {cbsa, title, kind}}.
    The header row is located by name; the build stops if a column is
    missing or a county appears in two CBSAs."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    header = None
    out: dict[str, dict] = {}
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if header is None:
            if all(h in cells for h in DELINEATION_HEADERS):
                header = {h: cells.index(h) for h in DELINEATION_HEADERS}
            continue
        code = cells[header["CBSA Code"]]
        if not re.fullmatch(r"\d{5}", code):
            continue  # footnotes
        st, co = cells[header["FIPS State Code"]], cells[header["FIPS County Code"]]
        if not (re.fullmatch(r"\d{1,2}", st) and re.fullmatch(r"\d{1,3}", co)):
            raise BuildError(f"delineation row with bad FIPS {st!r} {co!r}")
        fips = st.zfill(2) + co.zfill(3)
        kind_txt = cells[header["Metropolitan/Micropolitan Statistical Area"]].lower()
        kind = "metro" if kind_txt.startswith("metropolitan") else "micro" if kind_txt.startswith("micropolitan") else None
        if kind is None:
            raise BuildError(f"delineation: unknown CBSA type {kind_txt!r}")
        if fips in out and out[fips]["cbsa"] != code:
            raise BuildError(f"county {fips} in two CBSAs")
        out[fips] = {"cbsa": code, "title": cells[header["CBSA Title"]], "kind": kind}
    if header is None:
        raise BuildError(f"{path}: header row with {DELINEATION_HEADERS} not found")
    return out


def build_geographies(geos: dict, county_keys, state_keys, cbsa_keys, gaz: dict, delin: dict) -> dict:
    states = []
    state_name = {}
    for st in sorted(state_keys):
        g = geos.get(f"0400000US{st}")
        if not g:
            raise BuildError(f"state {st} missing from the geography file")
        state_name[st] = g["name"]
        states.append({"id": st, "name": g["name"], "abbr": g["stusab"]})
    cbsa_name = {}
    for cb in cbsa_keys:
        g = geos.get(f"310M700US{cb}") or geos.get(f"310M800US{cb}")
        cbsa_name[cb] = g["name"] if g else None
    counties = []
    for fips in sorted(county_keys):
        g = geos.get(f"0500000US{fips}")
        if not g:
            raise BuildError(f"county {fips} missing from the geography file")
        name = g["name"].rsplit(", ", 1)[0]
        d = delin.get(fips)
        cb = d["cbsa"] if d and d["cbsa"] in cbsa_keys else None
        gz = gaz.get(fips, {})
        counties.append({
            "id": fips, "name": name, "state": fips[:2], "state_name": state_name.get(fips[:2]),
            "county": fips[2:], "cbsa": cb, "cbsa_name": cbsa_name.get(cb) if cb else None,
            "lat": gz.get("lat"), "lon": gz.get("lon"), "land_sqmi": gz.get("land_sqmi"),
        })
    members: dict[str, list] = defaultdict(list)
    kinds = {}
    for fips, d in delin.items():
        members[d["cbsa"]].append(fips)
        kinds[d["cbsa"]] = d["kind"]
    cbsas = []
    for cb in sorted(cbsa_keys):
        if cb not in kinds:
            raise BuildError(f"ACS CBSA {cb} is not in the delineation file")
        cbsas.append({"id": cb, "name": cbsa_name[cb], "kind": kinds[cb], "counties": sorted(members[cb])})
    return {
        "vintage": f"{ACS_PERIOD} ACS 5-year geography; OMB July 2023 CBSA delineation; {ACS_YEAR} Gazetteer",
        "counties": counties, "states": states, "cbsas": cbsas,
    }


# ------------------------------------------------------ 2020 Religion Census

CATHOLIC_CODE = "081"
CATHOLIC_NAME = "Catholic Church"


def _fips(v, width: int, where: str) -> str | None:
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        raise BuildError(f"{where}: FIPS {v!r} stored as a number (leading zeros lost?)")
    s = str(v).strip()
    if re.fullmatch(rf"\d{{{width}}}", s):
        return s
    return None


def _num(v, where: str):
    if v is None or v == "":
        return None
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise BuildError(f"{where}: non-numeric value {v!r}")
    if v < 0:
        raise BuildError(f"{where}: negative count {v!r}")
    return int(v) if float(v).is_integer() else v


def _rows_with_header(rows, required: list[str], where: str):
    rows = list(rows)
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in r]
        if all(h in cells for h in required):
            ix = {h: cells.index(h) for h in cells if h}
            return ix, rows[i + 1:]
    raise BuildError(f"{where}: header with {required} not found")


def parse_usrc_summary(rows, key_col: str, width: int, where: str) -> dict[str, dict]:
    """Summary sheet -> {fips: {'population_2020', 'name'}}; 'Totals' rows are
    skipped by name, any other non-FIPS row with data stops the build."""
    ix, body = _rows_with_header(rows, [key_col, "2020 Population"], where)
    out: dict[str, dict] = {}
    for r in body:
        k = r[ix[key_col]] if ix[key_col] < len(r) else None
        if k is None and all(c is None for c in r):
            continue
        f = _fips(k, width, where)
        if f is None:
            if str(k).strip().lower() in ("total", "totals"):
                continue
            raise BuildError(f"{where}: unexpected key {k!r}")
        if f in out:
            raise BuildError(f"{where}: duplicate row for {f}")
        name_col = next((c for c in ("County Name", "Metro Name", "State Name") if c in ix), None)
        out[f] = {"population_2020": _num(r[ix["2020 Population"]], f"{where} {f}"),
                  "name": r[ix[name_col]] if name_col else None,
                  "state_name": r[ix["State Name"]] if "State Name" in ix else None}
    return out


def parse_usrc_group(rows, key_col: str, width: int, where: str) -> dict[str, dict]:
    """Group-detail sheet -> {fips: {'adherents', 'congregations'}} for the
    Catholic Church. Rejects: a FIPS listed twice for the group, a group code
    whose name is not 'Catholic Church', a 'Catholic Church' row under another
    code, and keys that are not FIPS codes."""
    ix, body = _rows_with_header(rows, [key_col, "Group Code", "Group Name", "Congregations", "Adherents"], where)
    out: dict[str, dict] = {}
    for r in body:
        if all(c is None for c in r):
            continue
        code = str(r[ix["Group Code"]]).strip() if r[ix["Group Code"]] is not None else None
        gname = str(r[ix["Group Name"]]).strip() if r[ix["Group Name"]] is not None else None
        if code != CATHOLIC_CODE and gname != CATHOLIC_NAME:
            continue
        if code != CATHOLIC_CODE or gname != CATHOLIC_NAME:
            raise BuildError(f"{where}: ambiguous Catholic row: code {code!r} name {gname!r}")
        k = r[ix[key_col]]
        f = _fips(k, width, where)
        if f is None:
            raise BuildError(f"{where}: Catholic row with key {k!r} that is not a FIPS/CBSA code")
        if f in out:
            raise BuildError(f"{where}: duplicate Catholic Church rows for {f}")
        out[f] = {"adherents": _num(r[ix["Adherents"]], f"{where} {f}"),
                  "congregations": _num(r[ix["Congregations"]], f"{where} {f}"),
                  "name": r[ix["County Name"]] if "County Name" in ix else None,
                  "state_name": r[ix["State Name"]] if "State Name" in ix else None}
    return out


def check_usrc_consistency(group: dict, summary: dict, where: str) -> None:
    for f, g in group.items():
        s = summary.get(f)
        if s is None:
            raise BuildError(f"{where}: Catholic row for {f} has no summary row")
        for k in ("name", "state_name"):
            if g.get(k) and s.get(k) and str(g[k]).strip() != str(s[k]).strip():
                raise BuildError(f"{where}: {f} is {g[k]!r} in one sheet and {s[k]!r} in the other")


def load_usrc(raw: Path) -> dict | None:
    files = Path(raw) / "usrc2020" / "files"
    det, summ = files / "2020_USRC_Group_Detail.xlsx", files / "2020_USRC_Summaries.xlsx"
    if not (det.exists() and summ.exists()):
        return None
    import openpyxl
    wd = openpyxl.load_workbook(det, read_only=True, data_only=True)
    ws_ = openpyxl.load_workbook(summ, read_only=True, data_only=True)

    def sheet(wb, name):
        if name not in wb.sheetnames:
            raise BuildError(f"Religion Census sheet {name!r} not found; sheets are {wb.sheetnames}")
        return list(wb[name].iter_rows(values_only=True))

    out = {
        "county_summary": parse_usrc_summary(sheet(ws_, "2020 County Summary"), "FIPS", 5, "USRC county summary"),
        "state_summary": parse_usrc_summary(sheet(ws_, "2020 State Summary"), "State Code", 2, "USRC state summary"),
        "metro_summary": parse_usrc_summary(sheet(ws_, "2020 Metro Summary"), "CBSA Code", 5, "USRC metro summary"),
        "county": parse_usrc_group(sheet(wd, "2020 Group by County"), "FIPS", 5, "USRC group by county"),
        "state": parse_usrc_group(sheet(wd, "2020 Group by State"), "StateCode", 2, "USRC group by state"),
        "metro": parse_usrc_group(sheet(wd, "2020 Group by Metro"), "CBSA Code", 5, "USRC group by metro"),
    }
    check_usrc_consistency(out["county"], out["county_summary"], "USRC county")
    check_usrc_consistency(out["state"], out["state_summary"], "USRC state")
    check_usrc_consistency(out["metro"], out["metro_summary"], "USRC metro")
    return out


CATHOLIC_METRICS = ("catholic_adherents", "catholic_congregations",
                    "catholic_adherents_per_1000", "catholic_congregations_per_100k")


def catholic_obs(group_row: dict | None, population_2020, *, flags_if_missing=("not_published",)) -> dict:
    """The four Catholic metrics for one unit. Denominator: the Religion
    Census's own 2020 population column (same year), all ages."""
    if group_row is None:
        return {m: obs(None, None, flags_if_missing) for m in CATHOLIC_METRICS}
    a, c = group_row["adherents"], group_row["congregations"]
    pop = obs(population_2020, None, [])
    o = {
        "catholic_adherents": obs(a, None, [] if a is not None else ["not_published"]),
        "catholic_congregations": obs(c, None, [] if c is not None else ["not_published"]),
    }
    o["catholic_adherents_per_1000"] = ratio_obs(o["catholic_adherents"], pop, 1000, subset=True,
                                                 extra_flags=["denominator_all_ages"])
    o["catholic_congregations_per_100k"] = ratio_obs(o["catholic_congregations"], pop, 100000, subset=False,
                                                     extra_flags=["denominator_all_ages"])
    return o


def join_usrc_counties(county_keys, usrc: dict | None) -> tuple[dict, dict]:
    """County Catholic observations keyed by ACS 2024 FIPS. Connecticut's
    2020 data are by its eight old counties while ACS 2024 uses planning
    regions, so those are null + geo_changed rather than forced together.
    Returns (observations, report)."""
    out = {}
    report = {"joined": 0, "no_catholic_row": [], "geo_changed": [], "not_covered": [],
              "usrc_unjoined": []}
    if usrc is None:
        return {k: {m: obs(None, None, ["not_ingested"]) for m in CATHOLIC_METRICS} for k in county_keys}, report
    keys = set(county_keys)
    for k in sorted(keys):
        summ = usrc["county_summary"].get(k)
        if summ is None:
            if k.startswith("09"):
                out[k] = catholic_obs(None, None, flags_if_missing=["geo_changed"])
                report["geo_changed"].append(k)
            else:
                out[k] = catholic_obs(None, None, flags_if_missing=["not_published"])
                report["not_covered"].append(k)
            continue
        row = usrc["county"].get(k)
        if row is None:
            report["no_catholic_row"].append(k)
        else:
            report["joined"] += 1
        out[k] = catholic_obs(row, summ["population_2020"])
    report["usrc_unjoined"] = sorted(set(usrc["county_summary"]) - keys)
    return out, report


def join_usrc_cbsas(cbsa_members: dict[str, list], usrc: dict | None) -> tuple[dict, dict]:
    """CBSA Catholic observations from the Religion Census's own metro rows.
    The 2020 study used an earlier CBSA delineation, so a metro's value is
    used only when its 2020 population equals the sum of the 2020 populations
    of the counties the 2023 delineation puts in it (same territory);
    otherwise null + geo_changed."""
    report = {"joined": 0, "geo_changed": []}
    if usrc is None:
        return {k: {m: obs(None, None, ["not_ingested"]) for m in CATHOLIC_METRICS} for k in cbsa_members}, report
    out = {}
    cs = usrc["county_summary"]
    for cb, members in cbsa_members.items():
        ms = usrc["metro_summary"].get(cb)
        same = (ms is not None and members and all(m in cs for m in members)
                and sum(cs[m]["population_2020"] for m in members) == ms["population_2020"])
        if not same:
            out[cb] = catholic_obs(None, None, flags_if_missing=["geo_changed"])
            report["geo_changed"].append(cb)
            continue
        out[cb] = catholic_obs(usrc["metro"].get(cb), ms["population_2020"])
        report["joined"] += 1
    return out, report


# ----------------------------------------------------------- public venues

def venue_counts(communities: dict | None, geographies: dict) -> tuple[dict, dict]:
    """Counts of communities.json rows per county / state / CBSA.
    Join on geo_id, else on (county_name, state abbreviation)."""
    report = {"rows": 0, "joined": 0, "unjoined": []}
    if communities is None:
        return None, report
    counties = {c["id"]: c for c in geographies["counties"]}
    abbr_to_state = {s["abbr"]: s["id"] for s in geographies["states"]}
    by_name = defaultdict(list)
    for c in geographies["counties"]:
        by_name[(c["name"].lower(), c["state"])].append(c["id"])
    county_n, county_unverified = Counter(), Counter()
    state_n, state_unverified = Counter(), Counter()
    for v in communities.get("venues", []):
        report["rows"] += 1
        unverified = v.get("verification") != "fetched_official_page"
        st = abbr_to_state.get((v.get("state") or "").upper())
        gid = v.get("geo_id")
        fips = None
        if gid and gid in counties:
            fips = gid
        elif v.get("county_name") and st:
            hits = by_name.get((v["county_name"].strip().lower(), st), [])
            fips = hits[0] if len(hits) == 1 else None
        if fips:
            report["joined"] += 1
            county_n[fips] += 1
            county_unverified[fips] += unverified
            st = fips[:2]
        else:
            report["unjoined"].append(v.get("id"))
        if st:
            state_n[st] += 1
            state_unverified[st] += unverified
    return {"county": (county_n, county_unverified), "state": (state_n, state_unverified)}, report


def venue_obs(n: int, unverified: int) -> dict:
    return obs(n, None, ["unvalidated"] if unverified else [])


# ------------------------------------------------------------- projections

def albers(lat1: float, lat2: float, lat0: float, lon0: float):
    """Spherical Albers equal-area conic; returns f(lon, lat) -> (x, y) with
    y pointing up (unit sphere)."""
    p1, p2, p0 = map(math.radians, (lat1, lat2, lat0))
    n = (math.sin(p1) + math.sin(p2)) / 2
    C = math.cos(p1) ** 2 + 2 * n * math.sin(p1)
    rho0 = math.sqrt(C - 2 * n * math.sin(p0)) / n

    def f(lon: float, lat: float):
        th = n * math.radians(lon - lon0)
        rho = math.sqrt(max(C - 2 * n * math.sin(math.radians(lat)), 0)) / n
        return rho * math.sin(th), rho0 - rho * math.cos(th)
    return f


def norm_lon(lon: float, state: str) -> float:
    # Alaska's Aleutians cross the antimeridian.
    return lon - 360 if state == "02" and lon > 0 else lon


def dp_simplify(pts: list, tol: float) -> list:
    """Douglas-Peucker on an open polyline (endpoints kept)."""
    if len(pts) < 3 or tol <= 0:
        return list(pts)
    keep = [False] * len(pts)
    keep[0] = keep[-1] = True
    stack = [(0, len(pts) - 1)]
    t2 = tol * tol
    while stack:
        a, b = stack.pop()
        ax, ay = pts[a]
        bx, by = pts[b]
        dx, dy = bx - ax, by - ay
        L = dx * dx + dy * dy
        best, bi = -1.0, -1
        for i in range(a + 1, b):
            px, py = pts[i]
            if L == 0:
                d = (px - ax) ** 2 + (py - ay) ** 2
            else:
                t = ((px - ax) * dx + (py - ay) * dy) / L
                t = 0 if t < 0 else 1 if t > 1 else t
                d = (px - ax - t * dx) ** 2 + (py - ay - t * dy) ** 2
            if d > best:
                best, bi = d, i
        if bi > 0 and best > t2:
            keep[bi] = True
            stack.append((a, bi))
            stack.append((bi, b))
    return [p for p, k in zip(pts, keep) if k]


class Topology:
    """County rings split into arcs at junctions, so a border shared by two
    counties is simplified once and stays shared (no slivers or gaps)."""

    def __init__(self, features: dict[str, list[list[tuple]]]):
        self.features = features  # fips -> rings (closed lon/lat lists)
        sharers: dict[tuple, set] = defaultdict(set)
        for fid, rings in features.items():
            for ring in rings:
                for p, q in zip(ring, ring[1:]):
                    sharers[(p, q) if p <= q else (q, p)].add(fid)
        self.sharers = {k: frozenset(v) for k, v in sharers.items()}
        self.ring_arcs: dict[str, list[list[tuple]]] = {}
        for fid, rings in features.items():
            self.ring_arcs[fid] = [self._split(r) for r in rings]

    def edge_sharers(self, p, q):
        return self.sharers[(p, q) if p <= q else (q, p)]

    def _split(self, ring):
        pts = ring[:-1] if ring[0] == ring[-1] else ring
        n = len(pts)
        if n < 3:
            return [(tuple(ring), frozenset())]
        es = [self.edge_sharers(pts[i], pts[(i + 1) % n]) for i in range(n)]
        junctions = [i for i in range(n) if es[i - 1] != es[i]]
        if not junctions:
            return [(tuple(pts + [pts[0]]), es[0])]
        arcs = []
        for j, start in enumerate(junctions):
            end = junctions[(j + 1) % len(junctions)]
            idx = list(range(start, end + 1)) if end > start else list(range(start, n)) + list(range(0, end + 1))
            arcs.append((tuple(pts[i % n] for i in idx), es[start]))
        return arcs


def fmt_tenths(v: int) -> str:
    """Integer tenths -> shortest decimal string (-15 -> '-1.5', 20 -> '2')."""
    a = abs(v)
    return ("-" if v < 0 else "") + str(a // 10) + ("" if a % 10 == 0 else "." + str(a % 10))


def ring_path(pts_tenths: list[tuple[int, int]]) -> str:
    """Closed ring of integer tenth-pixel points -> 'Mx,yl dx,dy ... z'."""
    out = ["M", fmt_tenths(pts_tenths[0][0]), ",", fmt_tenths(pts_tenths[0][1]), "l"]
    px, py = pts_tenths[0]
    first = True
    for x, y in pts_tenths[1:]:
        dx, dy = fmt_tenths(x - px), fmt_tenths(y - py)
        if not first and not dx.startswith("-"):
            out.append(" ")
        out.append(dx)
        if not dy.startswith("-"):
            out.append(",")
        out.append(dy)
        px, py = x, y
        first = False
    out.append("z")
    return "".join(out)


def clean_ring(pts: list[tuple[int, int]]) -> list[tuple[int, int]]:
    out = []
    for p in pts:
        if not out or out[-1] != p:
            out.append(p)
    if len(out) > 1 and out[0] == out[-1]:
        out.pop()
    # drop collinear-backtrack spikes created by rounding
    changed = True
    while changed and len(out) >= 3:
        changed = False
        for i in range(len(out)):
            if out[i - 1] == out[(i + 1) % len(out)]:
                del out[i]
                if i < len(out):
                    del out[i % len(out)]
                changed = True
                break
    return out


class Renderer:
    """Projects and simplifies through a Topology for one output file."""

    def __init__(self, topo: Topology, project, tol: float):
        self.topo, self.project, self.tol = topo, project, tol
        self.cache: dict[tuple, list] = {}

    def arc(self, fid: str, pts: tuple) -> list[tuple[int, int]]:
        key = pts if pts <= pts[::-1] else pts[::-1]
        if key not in self.cache:
            proj = [self.project(fid, lon, lat) for lon, lat in key]
            simp = dp_simplify(proj, self.tol)
            if key[0] == key[-1] and len(simp) < 4:  # closed arc: keep a triangle
                far = max(range(len(proj)), key=lambda i: (proj[i][0] - proj[0][0]) ** 2 + (proj[i][1] - proj[0][1]) ** 2)
                simp = dp_simplify(proj[:far + 1], self.tol) + dp_simplify(proj[far:], self.tol)[1:]
            self.cache[key] = [(round(x * 10), round(y * 10)) for x, y in simp]
        r = self.cache[key]
        return r if key == pts else r[::-1]

    def ring(self, fid: str, arcs) -> list[tuple[int, int]]:
        out: list = []
        for pts, _ in arcs:
            a = self.arc(fid, pts)
            out.extend(a if not out else a[1:])
        return clean_ring(out)

    def feature(self, fid: str) -> str:
        rings = [self.ring(fid, arcs) for arcs in self.topo.ring_arcs[fid]]
        good = [r for r in rings if len(r) >= 3]
        if not good:  # never drop a whole county: fall back to the unsimplified ring
            raw = self.topo.features[fid]
            biggest = max(raw, key=len)
            good = [clean_ring([(round(x * 10), round(y * 10))
                                for x, y in (self.project(fid, lo, la) for lo, la in biggest)])]
        return "".join(ring_path(r) for r in good if len(r) >= 3)

    def outline(self, fids: set[str]) -> str:
        """Outline of a group of counties: arcs whose sharers include exactly
        one member, chained into rings."""
        segs = []
        seen = set()
        for fid in sorted(fids):
            for arcs in self.topo.ring_arcs[fid]:
                for pts, sh in arcs:
                    if len(sh & fids) != 1:
                        continue
                    key = pts if pts <= pts[::-1] else pts[::-1]
                    if key in seen:
                        continue
                    seen.add(key)
                    segs.append(self.arc(fid, pts))
        rings = chain(segs)
        return "".join(ring_path(r) for r in (clean_ring(x) for x in rings) if len(r) >= 3)


def chain(segs: list[list[tuple]]) -> list[list[tuple]]:
    ends = defaultdict(list)
    for i, s in enumerate(segs):
        ends[s[0]].append(i)
        ends[s[-1]].append(i)
    used = [False] * len(segs)
    rings = []
    for i, s in enumerate(segs):
        if used[i]:
            continue
        used[i] = True
        ring = list(s)
        while ring[-1] != ring[0]:
            nxt = next((j for j in ends[ring[-1]] if not used[j]), None)
            if nxt is None:
                break
            used[nxt] = True
            t = segs[nxt]
            ring.extend(t[1:] if t[0] == ring[-1] else t[::-1][1:])
        rings.append(ring)
    return rings


def read_counties_shp(zip_path: Path) -> tuple[dict, dict]:
    """Cartographic boundary county shapefile -> ({fips: rings}, {fips: state})."""
    import shapefile
    z = zipfile.ZipFile(zip_path)
    base = next(n[:-4] for n in z.namelist() if n.endswith(".shp"))
    r = shapefile.Reader(shp=io.BytesIO(z.read(base + ".shp")), dbf=io.BytesIO(z.read(base + ".dbf")),
                         shx=io.BytesIO(z.read(base + ".shx")), encoding="utf-8")
    feats, states = {}, {}
    for sr in r.iterShapeRecords():
        fips = str(sr.record["GEOID"])
        if not re.fullmatch(r"\d{5}", fips):
            raise BuildError(f"shapefile GEOID {fips!r}")
        if fips in feats:
            raise BuildError(f"shapefile duplicate GEOID {fips}")
        s = sr.shape
        parts = list(s.parts) + [len(s.points)]
        st = fips[:2]
        rings = []
        for a, b in zip(parts, parts[1:]):
            ring = [(norm_lon(p[0], st), p[1]) for p in s.points[a:b]]
            if ring and ring[0] != ring[-1]:
                ring.append(ring[0])
            if len(ring) >= 4:
                rings.append(ring)
        feats[fips], states[fips] = rings, st
    return feats, states


# National layout: lower 48 in Albers (29.5, 45.5); Alaska, Hawaii and Puerto
# Rico as insets, each in its own Albers, scaled relative to the lower 48
# (Alaska 0.35 as in the usual "Albers USA" composite).
NATIONAL_W, NATIONAL_H, PAD = 975.0, 610.0, 10.0
INSETS = {
    "02": {"proj": (55, 65, 50, -154), "scale": 0.35, "anchor": "bottom-left"},
    "15": {"proj": (8, 18, 19.9, -157), "scale": 1.0, "anchor": "after-02"},
    "72": {"proj": (8, 18, 18.2, -66.4), "scale": 1.0, "anchor": "bottom-right"},
}
LOWER48 = (29.5, 45.5, 37.5, -96)


def _bbox(points):
    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    return min(xs), min(ys), max(xs), max(ys)


def national_projector(feats: dict, states: dict):
    """Return (project(fid, lon, lat) -> svg px, viewBox string, inset boxes)."""
    base = albers(*LOWER48)
    l48 = [base(lo, la) for f, rings in feats.items() if states[f] not in INSETS for r in rings for lo, la in r]
    x0, y0, x1, y1 = _bbox(l48)
    s = min((NATIONAL_W - 2 * PAD) / (x1 - x0), (NATIONAL_H - 2 * PAD) / (y1 - y0))
    tx = PAD + ((NATIONAL_W - 2 * PAD) - s * (x1 - x0)) / 2 - s * x0
    ty = PAD + s * y1  # svg y = ty - s*y

    transforms = {None: (base, s, tx, ty)}
    boxes = {}
    cursor_left = PAD
    for st in ("02", "15", "72"):
        cfg = INSETS[st]
        pts = [(lo, la) for f, rings in feats.items() if states[f] == st for r in rings for lo, la in r]
        if not pts:
            continue
        pr = albers(*cfg["proj"])
        pp = [pr(lo, la) for lo, la in pts]
        a0, b0, a1, b1 = _bbox(pp)
        k = s * cfg["scale"]
        w, h = k * (a1 - a0), k * (b1 - b0)
        bottom = NATIONAL_H - PAD
        if cfg["anchor"] == "bottom-right":
            left = NATIONAL_W - PAD - w
        else:
            left = cursor_left
            cursor_left = left + w + 2 * PAD
        itx = left - k * a0
        ity = bottom - h + k * b1
        transforms[st] = (pr, k, itx, ity)
        boxes[st] = [round(left, 1), round(bottom - h, 1), round(w, 1), round(h, 1)]

    def project(fid, lon, lat):
        pr, k, tx_, ty_ = transforms.get(states[fid], transforms[None])
        x, y = pr(lon, lat)
        return tx_ + k * x, ty_ - k * y
    return project, f"0 0 {int(NATIONAL_W)} {int(NATIONAL_H)}", boxes


def state_projector(feats: dict, fids: list[str], max_w=800.0, max_h=600.0):
    pts = [(lo, la) for f in fids for r in feats[f] for lo, la in r]
    lons, lats = [p[0] for p in pts], [p[1] for p in pts]
    la0, la1, lo0, lo1 = min(lats), max(lats), min(lons), max(lons)
    span = max(la1 - la0, 0.5)
    pr = albers(la0 + span / 6, la1 - span / 6, (la0 + la1) / 2, (lo0 + lo1) / 2)
    pp = [pr(lo, la) for lo, la in pts]
    x0, y0, x1, y1 = _bbox(pp)
    s = min((max_w - 2 * PAD) / (x1 - x0), (max_h - 2 * PAD) / (y1 - y0))
    W = math.ceil(s * (x1 - x0) + 2 * PAD)
    H = math.ceil(s * (y1 - y0) + 2 * PAD)

    def project(fid, lon, lat):
        x, y = pr(lon, lat)
        return PAD + s * (x - x0), PAD + s * (y1 - y)
    desc = (f"Albers equal-area conic, standard parallels {la0 + span / 6:.2f} and {la1 - span / 6:.2f}, "
            f"central meridian {(lo0 + lo1) / 2:.2f}, fitted to the state's bounding box")
    return project, f"0 0 {W} {H}", desc


def build_geo(shp_zip: Path, keep: set[str] | None = None, tol_national=0.1, tol_state=0.15):
    """Return {'us_counties': {...}, 'state_08': {...}, ...}."""
    feats, states = read_counties_shp(shp_zip)
    if keep is not None:
        feats = {k: v for k, v in feats.items() if k in keep}
        states = {k: v for k, v in states.items() if k in keep}
    topo = Topology(feats)
    out = {}
    project, vb, boxes = national_projector(feats, states)
    rnd = Renderer(topo, project, tol_national)
    by_state = defaultdict(set)
    for f, st in states.items():
        by_state[st].add(f)
    out["us_counties"] = {
        "viewBox": vb,
        "projection": ("Albers equal-area conic (standard parallels 29.5 and 45.5, central meridian -96) "
                       "for the lower 48; Alaska (scale 0.35), Hawaii and Puerto Rico as insets in their "
                       "own Albers projections"),
        "insets": boxes,
        "paths": {f: rnd.feature(f) for f in sorted(feats)},
        "states": {st: rnd.outline(fs) for st, fs in sorted(by_state.items())},
    }
    for st, fs in sorted(by_state.items()):
        sp, svb, desc = state_projector(feats, sorted(fs))
        r = Renderer(topo, sp, tol_state)
        out[f"state_{st}"] = {
            "viewBox": svb, "projection": desc,
            "paths": {f: r.feature(f) for f in sorted(fs)},
            "states": {st: r.outline(fs)},
        }
    return out


# ------------------------------------------------------------ sources/metrics

def manifest_source(manifest: dict, sid: str) -> dict:
    return manifest.get("sources", {}).get(sid, {})


def files_of(manifest: dict, sid: str, pattern: str = "") -> list[dict]:
    out = []
    for f in manifest_source(manifest, sid).get("files", []):
        if pattern and not re.search(pattern, f.get("path", "")):
            continue
        out.append({"url": f["url"], "sha256": f.get("sha256"), "path": f.get("path"),
                    "retrieved_at": f.get("retrieved_at")})
    return out


def latest(files: list[dict]):
    ts = [f.get("retrieved_at") for f in files if f.get("retrieved_at")]
    return max(ts) if ts else None


def base_sources(manifest: dict, status: dict) -> list[dict]:
    """One entry per SPEC section 4 source (plus the Census geography files)."""
    acs_files = files_of(manifest, "acs5", r"(sf/acsdt5y|sf/Geos|groups_)")
    shell_files = files_of(manifest, "acs5", r"Table_Shells")
    geo_files = files_of(manifest, "census_geo")
    usrc_files = files_of(manifest, "usrc2020", r"files/")
    usrc_pages = files_of(manifest, "usrc2020", r"pages/")

    def errors(sid):
        m = manifest_source(manifest, sid)
        return "; ".join(m.get("errors", []))[:400]

    S = []
    S.append({
        "id": "acs5", "title": "American Community Survey 5-year estimates, 2020-2024 (detailed tables B01001, B01003, B12002)",
        "provider": "U.S. Census Bureau",
        "landing_url": "https://www.census.gov/data/developers/data-sets/acs-5year/2024.html",
        "dataset_version": "2024 ACS 5-year, table-based Summary File",
        "reference_period": ACS_PERIOD, "retrieved_at": latest(acs_files),
        "license": "U.S. Government work; public domain in the United States (17 U.S.C. 105). Census Bureau asks that the source be cited.",
        "access": "public download (table-based Summary File; no key needed)",
        "citation": "U.S. Census Bureau, 2020-2024 American Community Survey 5-Year Estimates, Tables B01001, B01003 and B12002 (table-based Summary File).",
        "status": status.get("acs5", "not_ingested"), "files": [{"url": f["url"], "sha256": f["sha256"]} for f in acs_files],
        "notes": ("Summary File .dat and geography files are filtered by pipeline/fetch.py to the nation, state, county "
                  "and CBSA rows before they are stored; the sha256 is of the filtered file. Cell codes are verified "
                  "against the API group metadata labels and the table shells at build time."),
    })
    S.append({
        "id": "acs5_variables", "title": "ACS 2024 5-year variable documentation (group metadata and table shells)",
        "provider": "U.S. Census Bureau",
        "landing_url": "https://api.census.gov/data/2024/acs/acs5/variables.html",
        "dataset_version": "2024 ACS 5-year", "reference_period": ACS_PERIOD,
        "retrieved_at": latest(shell_files + [f for f in acs_files if "groups_" in (f.get("path") or "")]),
        "license": "U.S. Government work; public domain in the United States.",
        "access": "public (API metadata JSON, no key needed; Summary File documentation)",
        "citation": "U.S. Census Bureau, American Community Survey 2024 5-year API variables and table shells.",
        "status": status.get("acs5_variables", "not_ingested"),
        "files": [{"url": f["url"], "sha256": f["sha256"]} for f in shell_files + [f for f in acs_files if "groups_" in (f.get("path") or "")]],
        "notes": "Used only to verify that each cell code carries the expected label; the build fails on a mismatch.",
    })
    S.append({
        "id": "census_geo", "title": "Census geography: 2024 county Gazetteer, 2024 cartographic boundary counties (1:20,000,000), July 2023 CBSA delineation",
        "provider": "U.S. Census Bureau; Office of Management and Budget (delineation)",
        "landing_url": "https://www.census.gov/geographies/reference-files.html",
        "dataset_version": "Gazetteer 2024; cb_2024_us_county_20m; list1_2023 (OMB Bulletin 23-01)",
        "reference_period": "2024 boundaries; July 2023 delineation", "retrieved_at": latest(geo_files),
        "license": "U.S. Government work; public domain in the United States.",
        "access": "public download",
        "citation": "U.S. Census Bureau, 2024 Gazetteer Files; 2024 Cartographic Boundary Files (counties, 1:20,000,000); Core Based Statistical Areas delineation file, July 2023.",
        "status": status.get("census_geo", "not_ingested"),
        "files": [{"url": f["url"], "sha256": f["sha256"]} for f in geo_files],
        "notes": "County internal points (lat/lon), land area, county boundaries for the maps, and county-to-CBSA membership.",
    })
    S.append({
        "id": "usrc2020", "title": "2020 U.S. Religion Census: Religious Congregations & Adherents Study",
        "provider": "Association of Statisticians of American Religious Bodies (ASARB)",
        "landing_url": "https://www.usreligioncensus.org/node/1639",
        "dataset_version": "2020 summary and group detail workbooks (files updated 2023-06-23 per the landing page)",
        "reference_period": "2020", "retrieved_at": latest(usrc_files + usrc_pages),
        "license": ("Workbooks carry 'Copyright (c) 2022 by the Association of Statisticians of American Religious Bodies "
                    "(ASARB). All rights reserved' and a suggested citation. The landing page offers them as free "
                    "downloads and states no further terms of use; the ARDA archives the same study with a DOI and "
                    "citation. Only per-county Catholic Church counts and ratios derived from them are republished "
                    "here, with the citation."),
        "access": "public download (XLSX)",
        "citation": ("Clifford Grammich, Erica Dollhopf, Mary Gautier, Richard Houseal, Dale E. Jones, Alexei Krindatch, "
                     "Richie Stanley, and Scott Thumma. 2022. 2020 U.S. Religion Census: Religious Congregations & "
                     "Adherents Study. Association of Statisticians of American Religious Bodies. County file archived "
                     "at the ARDA: DOI 10.17605/OSF.IO/ET2A5."),
        "status": status.get("usrc2020", "not_ingested"),
        "files": [{"url": f["url"], "sha256": f["sha256"]} for f in usrc_files],
        "notes": ("Catholic Church = group code 081. Adherents are all ages and include people the Church counts as "
                  "members; they are not Mass attendees and not women 20-29. Connecticut is reported by its eight "
                  "former counties, which do not match the planning regions ACS 2024 uses. Puerto Rico is not covered."),
    })
    gated = [
        ("issp2022", "ISSP 2022: Family and Changing Gender Roles V", "GESIS - Leibniz Institute for the Social Sciences",
         "https://www.gesis.org/en/issp/data-and-documentation/family-and-changing-gender-roles/2022", "2022 (fieldwork varies by country)",
         "Registration required at GESIS; respondent-level files may not be redistributed."),
        ("issp_family_waves", "ISSP Family and Changing Gender Roles, earlier waves", "GESIS",
         "https://www.gesis.org/en/issp/data-and-documentation/family-and-changing-gender-roles", "1988, 1994, 2002, 2012",
         "Registration required at GESIS."),
        ("evs_wvs_2017_2022", "Joint EVS/WVS 2017-2022 dataset", "GESIS / World Values Survey Association",
         "https://www.gesis.org/en/european-values-study/data-and-documentation/joint-evs/wvs-2017-2022-dataset", "2017-2022",
         "Registration/terms acceptance required."),
        ("latinobarometro2024", "Latinobarómetro 2024", "Corporación Latinobarómetro",
         "https://www.latinobarometro.org/latinobarometro-2024", "2024", "Download requires registration on the publisher's site."),
        ("pew_datasets", "Pew Research Center downloadable datasets", "Pew Research Center",
         "https://www.pewresearch.org/datasets/", "varies", "Pew account required."),
        ("dhs", "Demographic and Health Surveys", "The DHS Program (ICF / USAID)",
         "https://dhsprogram.com/data/", "varies by country wave", "Account and project approval required."),
        ("prri_ava", "PRRI American Values Atlas", "PRRI", "https://ava.prri.org/religion/", "varies",
         "Interactive atlas; no bulk file download found; treated as gated."),
    ]
    for sid, title, prov, url, period, acc in gated:
        S.append({"id": sid, "title": title, "provider": prov, "landing_url": url, "dataset_version": None,
                  "reference_period": period, "retrieved_at": None, "license": "See provider terms; not redistributed.",
                  "access": acc, "citation": f"{prov}. {title}.", "status": "gated", "files": [],
                  "notes": "Importer only: a user-supplied licensed file can be loaded locally; nothing is shown until then."})
    simple = [
        ("gss", "General Social Survey", "NORC at the University of Chicago", "https://gssdataexplorer.norc.org/gss_data", "1972-2024"),
        ("nsfg", "National Survey of Family Growth 2022-2023 public-use files", "CDC National Center for Health Statistics",
         "https://www.cdc.gov/nchs/nsfg/nsfg-2022-2023-puf.htm", "2022-2023"),
        ("wpp2024", "UN World Population Prospects 2024", "United Nations, Department of Economic and Social Affairs, Population Division",
         "https://population.un.org/wpp/", "1950-2100 (estimates to 2023, projections after)"),
        ("worldbank", "World Bank World Development Indicators (fertility, population)", "World Bank",
         "https://data.worldbank.org/", "1990-2025"),
        ("eurostat", "Eurostat divorce indicators demo_ndivind", "Eurostat",
         "https://ec.europa.eu/eurostat/databrowser/view/demo_ndivind/default/table?lang=en", "varies"),
    ]
    mid = {"gss": "gss", "nsfg": "nsfg", "wpp2024": "wpp2024", "worldbank": "worldbank", "eurostat": "eurostat"}
    for sid, title, prov, url, period in simple:
        m = manifest_source(manifest, mid[sid])
        fl = files_of(manifest, mid[sid])
        st = "failed" if m.get("status") == "failed" and not fl else "not_ingested"
        S.append({"id": sid, "title": title, "provider": prov, "landing_url": url, "dataset_version": None,
                  "reference_period": period, "retrieved_at": latest(fl), "license": "See provider terms.",
                  "access": "public download", "citation": f"{prov}. {title}.", "status": st,
                  "files": [{"url": f["url"], "sha256": f["sha256"]} for f in fl],
                  "notes": (f"Fetch failed: {errors(mid[sid])}" if st == "failed" else
                            "Raw files fetched; aggregates are built by the survey/international pipeline, which "
                            "replaces this entry through pipeline/registry/ when it ingests them.")})
    mit = manifest_source(manifest, "mit_election")
    mit_files = [f for f in files_of(manifest, "mit_election") if not f["path"].endswith("dataset.json")]
    S.append({"id": "mit_election", "title": "County Presidential Election Returns 2000-2024",
              "provider": "MIT Election Data and Science Lab (Harvard Dataverse)",
              "landing_url": "https://doi.org/10.7910/DVN/VOQCHQ", "dataset_version": None, "reference_period": "2000-2024",
              "retrieved_at": latest(mit_files), "license": "See Dataverse terms (CC0 per dataset page, to verify on ingest).",
              "access": "public download", "citation": "MIT Election Data and Science Lab. County Presidential Election Returns 2000-2024. Harvard Dataverse, doi:10.7910/DVN/VOQCHQ.",
              "status": status.get("mit_election", "failed" if mit.get("status") == "failed" else "not_ingested"),
              "files": [{"url": f["url"], "sha256": f["sha256"]} for f in mit_files],
              "notes": f"Data file download failed: {errors('mit_election')}" if mit.get("status") == "failed" else ""})
    hs = manifest_source(manifest, "homeschool")
    S.append({"id": "jhu_homeschool_hub", "title": "Johns Hopkins Homeschool Hub state statistics",
              "provider": "Johns Hopkins University School of Education",
              "landing_url": "https://education.jhu.edu/edpolicy/policy-research-initiatives/homeschool-hub/",
              "dataset_version": None, "reference_period": "2025-26 release (partial)", "retrieved_at": None,
              "license": "See provider terms.", "access": "public web page",
              "citation": "Johns Hopkins Institute for Education Policy, Homeschool Hub.",
              "status": "failed" if hs.get("status") == "failed" else "not_ingested", "files": [],
              "notes": f"Fetch failed: {errors('homeschool')}" if hs.get("status") == "failed" else ""})
    S.append({"id": "communities", "title": "Raphael's Map community listings (data/communities.json)",
              "provider": "Raphael's Map (maintained listing of publicly advertised venues)",
              "landing_url": "communities.json", "dataset_version": None, "reference_period": "see verified_at per row",
              "retrieved_at": None, "license": "Listing compiled for this project; each row links to its public source.",
              "access": "in repository", "citation": "Raphael's Map community listings.",
              "status": status.get("communities", "not_ingested"), "files": [],
              "notes": "Public venues and organizations only, never private people. Rows marked listed_from_search are not verified."})
    return S


METRIC_BASE = {
    "source_id": "acs5", "class": "observed", "higher_is": "more", "unit": "persons",
}
ACS_CAVEAT = ("American Community Survey 5-year estimate pooled over 2020-2024 (not a 2024 point estimate); "
              "sample-based, with a 90% margin of error. ")


def metric_defs(resolved: dict | None) -> list[dict]:
    cells = (resolved or {}).get("cells", {})

    def ids(metric, fallback):
        return cells.get(metric, fallback)

    M = []

    def add(metric_id, label, short, universe, definition, method, src_ids, caveats, **kw):
        d = dict(METRIC_BASE)
        d.update({"metric_id": metric_id, "label": label, "short": short, "universe": universe,
                  "definition": definition, "source_variable_ids": src_ids, "method": method, "caveats": caveats})
        d.update(kw)
        M.append(d)

    add("total_population", "Total population", "Population", "Total population (all ages, both sexes)",
        "ACS B01003_001E: Total population.", "Published ACS cell.", ids("total_population", ["B01003_001E"]),
        ACS_CAVEAT + "County totals are controlled to the official population estimates, so the ACS publishes no "
        "sampling margin of error for them (shown as 0).")
    add("female_population", "Female population, all ages", "Female pop.", "Female residents, all ages",
        "ACS B01001_026E (Total: Female:).", "Published ACS cell.", ids("female_population", ["B01001_026E"]), ACS_CAVEAT)
    add("women_18_19", "Women aged 18-19", "Women 18-19", "Female residents aged 18 and 19",
        "ACS B01001_031E (Female: 18 and 19 years).", "Published ACS cell.", ids("women_18_19", ["B01001_031E"]), ACS_CAVEAT)
    add("women_20_24", "Women aged 20-24", "Women 20-24", "Female residents aged 20 to 24, any marital status",
        "ACS B01001_032E + B01001_033E + B01001_034E (Female: 20 years; 21 years; 22 to 24 years).",
        "Sum of published ACS age cells; margin of error sqrt(sum of squared cell MOEs), the Census Bureau's approximation.",
        ids("women_20_24", ["B01001_032E", "B01001_033E", "B01001_034E"]),
        ACS_CAVEAT + "All women of that age, married or not.")
    add("women_25_29", "Women aged 25-29", "Women 25-29", "Female residents aged 25 to 29, any marital status",
        "ACS B01001_035E (Female: 25 to 29 years).", "Published ACS cell.", ids("women_25_29", ["B01001_035E"]), ACS_CAVEAT)
    add("women_30_34", "Women aged 30-34", "Women 30-34", "Female residents aged 30 to 34, any marital status",
        "ACS B01001_036E (Female: 30 to 34 years).", "Published ACS cell.", ids("women_30_34", ["B01001_036E"]), ACS_CAVEAT)
    add("women_20_29", "Women aged 20-29", "Women 20-29", "Female residents aged 20 to 29, any marital status",
        "ACS B01001_032E + _033E + _034E + _035E.", "Sum of published ACS age cells; approximate margin of error.",
        ids("women_20_29", ["B01001_032E", "B01001_033E", "B01001_034E", "B01001_035E"]), ACS_CAVEAT)
    for key, age in B12002_NEVER_MARRIED_AGES.items():
        a, b = key.split("_")
        add(f"never_married_women_{key}", f"Never-married women aged {a}-{b}", f"Never-married women {a}-{b}",
            f"Female residents aged {a} to {b} who have never married (population 15 years and over)",
            f"ACS B12002 cell 'Female: Never married: {age}', found by its published label.",
            "Published ACS cell, located by label at build time.",
            ids(f"never_married_women_{key}", []),
            ACS_CAVEAT + "Never married is not the same as single: it includes people living with a partner and "
            "excludes the divorced and widowed. B12002 publishes 20-24 as one bin; no finer age split exists.")
    add("never_married_women_20_29", "Never-married women aged 20-29", "Never-married women 20-29",
        "Female residents aged 20 to 29 who have never married",
        "Sum of the B12002 never-married female cells for 20-24 and 25-29.",
        "Sum of published ACS cells; approximate margin of error.", ids("never_married_women_20_29", []),
        ACS_CAVEAT + "Never married is not the same as single.")
    add("never_married_share_women_20_29", "Share of women 20-29 who have never married", "Never-married share 20-29",
        "Female residents aged 20 to 29", "never_married_women_20_29 / women_20_29 x 100.",
        "Ratio of ACS estimates; margin of error by the Census Bureau's formula for a derived proportion (approximate). "
        "The build checks that B01001 and B12002 give identical female 20-29 totals at every geography.",
        ids("never_married_women_20_29", []) + ids("women_20_29", []),
        ACS_CAVEAT + "Derived from two published tables.", unit="percent", higher_is="more")
    add("women_20_29_per_1000", "Women aged 20-29 per 1,000 residents", "Women 20-29 per 1,000",
        "Total population (denominator)", "women_20_29 / total_population x 1,000.",
        "Ratio of ACS estimates; approximate margin of error.",
        ids("women_20_29", []) + ids("total_population", []), ACS_CAVEAT + "A density, not a count: small counties can rank high.",
        unit="per 1,000 residents")
    usrc_cav = ("2020 U.S. Religion Census. Adherents are people of all ages the Catholic Church counts as members "
                "(including children); not Mass attendance, not women 20-29, and a 2020 count shown next to "
                "2020-2024 ACS estimates. Connecticut's 2020 data are by its former counties and cannot be joined to "
                "the 2024 planning regions (flag geo_changed). Puerto Rico is not covered. A county with no Catholic "
                "Church row in the Group Detail file is shown as unavailable (not_published), not as zero.")
    add("catholic_adherents", "Catholic adherents (all ages)", "Catholic adherents",
        "Adherents of the Catholic Church, all ages, as reported to the 2020 U.S. Religion Census",
        "Group Detail workbook, sheet '2020 Group by County', Group Code 081 'Catholic Church', column Adherents.",
        "Published count.", ["USRC2020:GroupByCounty:081:Adherents"], usrc_cav, source_id="usrc2020", class_="proxy")
    add("catholic_congregations", "Catholic congregations", "Catholic congregations",
        "Catholic Church congregations (parishes and missions as reported)",
        "Group Detail workbook, '2020 Group by County', Group Code 081, column Congregations.",
        "Published count.", ["USRC2020:GroupByCounty:081:Congregations"], usrc_cav, source_id="usrc2020", class_="proxy")
    add("catholic_adherents_per_1000", "Catholic adherents per 1,000 residents (all ages)", "Catholics per 1,000",
        "Total 2020 population, all ages (denominator)",
        "catholic_adherents / Religion Census 2020 population x 1,000.",
        "Ratio; the denominator is the Religion Census's own '2020 Population' column (same year, same geography). "
        "Flagged denominator_all_ages.",
        ["USRC2020:GroupByCounty:081:Adherents", "USRC2020:CountySummary:2020 Population"], usrc_cav,
        source_id="usrc2020", class_="proxy", unit="per 1,000 residents")
    add("catholic_congregations_per_100k", "Catholic congregations per 100,000 residents", "Congregations per 100k",
        "Total 2020 population, all ages (denominator)",
        "catholic_congregations / Religion Census 2020 population x 100,000.",
        "Ratio against the Religion Census's own 2020 population column.",
        ["USRC2020:GroupByCounty:081:Congregations", "USRC2020:CountySummary:2020 Population"], usrc_cav,
        source_id="usrc2020", class_="proxy", unit="per 100,000 residents")
    add("rep_pres_share_2024", "Republican share of the 2024 presidential vote", "Rep. vote 2024",
        "Votes cast for president in the county, 2024",
        "Republican candidate votes / total presidential votes (MIT Election Lab county returns).",
        "Not computed: the source file has not been downloaded.", [],
        "Exploratory ecological proxy, untested against any individual-level outcome; off in the default weights. "
        "A county's vote says nothing about any individual.",
        source_id="mit_election", class_="unvalidated_hypothesis", unit="percent", higher_is="more", status="not_ingested")
    add("public_venues", "Listed public Catholic young-adult venues", "Public venues",
        "Rows of communities.json whose location joins to the county",
        "Count of listed public venues (parish young-adult groups, campus ministries, etc.) in the county.",
        "Count of rows; joined on geo_id or (county name, state).", [],
        "A curated, incomplete listing of public organizations; a listing does not measure how many people attend. "
        "Flag unvalidated when any counted row was listed from search rather than checked on its official page. "
        "Diocese-wide offices without a county are counted only at state level.",
        source_id="communities", class_="user_listing", unit="venues")
    for d in M:
        if "class_" in d:
            d["class"] = d.pop("class_")
    return M


def merge_registry(base_sources: list, base_metrics: list, registry_dir: Path) -> tuple[list, list, list]:
    """Union of the base registry and pipeline/registry/*.json fragments
    ({sources:[...], metrics:[...]}). A fragment entry replaces a base entry
    with the same id; two fragments disagreeing on one id stop the build."""
    notes = []
    src = {s["id"]: s for s in base_sources}
    met = {m["metric_id"]: m for m in base_metrics}
    from_frag_s, from_frag_m = {}, {}
    for p in sorted(Path(registry_dir).glob("*.json")) if Path(registry_dir).exists() else []:
        try:
            frag = json.loads(p.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            raise BuildError(f"{p}: not valid JSON ({e})")
        if not isinstance(frag, dict) or not ({"sources", "metrics"} & set(frag)):
            notes.append(f"{p.name} skipped: not a {{sources, metrics}} fragment")
            continue
        for s in frag.get("sources", []):
            sid = s.get("id")
            if not sid:
                raise BuildError(f"{p}: source without id")
            if sid in from_frag_s and from_frag_s[sid][1] != s:
                raise BuildError(f"source {sid} defined differently in {from_frag_s[sid][0]} and {p.name}")
            from_frag_s[sid] = (p.name, s)
            if sid in src:
                notes.append(f"source {sid} replaced by {p.name}")
            src[sid] = s
        for m in frag.get("metrics", []):
            mid = m.get("metric_id")
            if not mid:
                raise BuildError(f"{p}: metric without metric_id")
            if mid in from_frag_m and from_frag_m[mid][1] != m:
                raise BuildError(f"metric {mid} defined differently in {from_frag_m[mid][0]} and {p.name}")
            from_frag_m[mid] = (p.name, m)
            if mid in met:
                notes.append(f"metric {mid} replaced by {p.name}")
            met[mid] = m
    return list(src.values()), list(met.values()), notes


# ------------------------------------------------------------------- checks

def internal_checks(acs: dict, usrc: dict | None, geographies: dict) -> dict:
    """Consistency checks that use only the files themselves."""
    checks = {}
    # county sums vs state rows, state sums vs nation row (exact for controlled
    # cells, close otherwise).
    for metric in ("total_population", "women_20_24", "women_25_29", "never_married_women_20_24"):
        worst = (0.0, None)
        for st, so in acs["state"].items():
            vals = [o[metric]["e"] for k, o in acs["county"].items() if k[:2] == st]
            if not vals or any(v is None for v in vals) or so[metric]["e"] in (None, 0):
                continue
            rel = abs(sum(vals) - so[metric]["e"]) / so[metric]["e"]
            if rel > worst[0]:
                worst = (rel, st)
        nat = acs["us"].get("US", {}).get(metric, {}).get("e")
        st_sum = sum(o[metric]["e"] for k, o in acs["state"].items() if k != "72" and o[metric]["e"] is not None)
        checks[f"{metric}_county_sum_vs_state_max_rel_diff"] = {"value": round(worst[0], 6), "state": worst[1]}
        checks[f"{metric}_50_states_dc_sum_vs_us_row"] = {"sum": st_sum, "us_row": nat,
                                                          "rel_diff": round(abs(st_sum - nat) / nat, 6) if nat else None}
    if usrc:
        worst = (0.0, None)
        for st, row in usrc["state"].items():
            tot = sum(v["adherents"] or 0 for k, v in usrc["county"].items() if k[:2] == st)
            if row["adherents"]:
                rel = abs(tot - row["adherents"]) / row["adherents"]
                if rel > worst[0]:
                    worst = (rel, st)
        checks["catholic_adherents_county_sum_vs_state_max_rel_diff"] = {"value": round(worst[0], 6), "state": worst[1]}
        checks["usrc_catholic_county_rows"] = len(usrc["county"])
    checks["counties"] = len(geographies["counties"])
    checks["cbsas"] = len(geographies["cbsas"])
    return checks


# --------------------------------------------------------------------- main

COUNTY_METRICS = [
    "total_population", "female_population", "women_18_19", "women_20_24", "women_25_29", "women_30_34",
    "women_20_29", "never_married_women_20_24", "never_married_women_25_29", "never_married_women_20_29",
    "never_married_women_30_34", "never_married_share_women_20_29", "women_20_29_per_1000",
    "catholic_adherents", "catholic_congregations", "catholic_adherents_per_1000",
    "catholic_congregations_per_100k", "rep_pres_share_2024", "public_venues",
]


def dump(path: Path, obj) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(obj, ensure_ascii=False, separators=(",", ":"))
    path.write_text(text + "\n", encoding="utf-8")
    return len(text) + 1


def build(raw: Path, out: Path, registry_dir: Path = HERE / "registry", *, geo: bool = True, log=print) -> dict:
    raw, out = Path(raw), Path(out)
    manifest = json.loads((raw / "manifest.json").read_text(encoding="utf-8")) if (raw / "manifest.json").exists() else {}
    if manifest.get("synthetic") and out.resolve() == (HERE.parent / "data").resolve():
        raise BuildError("refusing to write synthetic fixture data into the published data/ directory")
    meta = raw / "acs5" / str(ACS_YEAR)
    acs, resolved = acs_observations(meta / "sf", meta)
    log(f"ACS: {len(acs['county'])} counties, {len(acs['state'])} states, {len(acs['cbsa'])} CBSAs, "
        f"nation={'US' in acs['us']}")
    geos = read_geos(meta / "sf" / f"Geos{ACS_YEAR}5YR.txt")
    gz = list((raw / "census_geo").glob("*_Gaz_counties_national.zip"))
    gaz = read_gazetteer(gz[0]) if gz else {}
    delin_p = raw / "census_geo" / "list1_2023.xlsx"
    if not delin_p.exists():
        raise BuildError(f"missing CBSA delineation {delin_p}")
    delin = read_delineation(delin_p)
    geographies = build_geographies(geos, acs["county"].keys(), acs["state"].keys(), acs["cbsa"].keys(), gaz, delin)
    missing_gaz = [c["id"] for c in geographies["counties"] if c["lat"] is None]

    usrc = load_usrc(raw)
    c_cath, c_rep = join_usrc_counties(acs["county"].keys(), usrc)
    members = {c["id"]: c["counties"] for c in geographies["cbsas"]}
    b_cath, b_rep = join_usrc_cbsas(members, usrc)
    s_cath = {}
    for st in acs["state"]:
        if usrc is None:
            s_cath[st] = {m: obs(None, None, ["not_ingested"]) for m in CATHOLIC_METRICS}
        else:
            ss = usrc["state_summary"].get(st)
            s_cath[st] = catholic_obs(usrc["state"].get(st), ss["population_2020"] if ss else None) if ss \
                else catholic_obs(None, None)

    comm_p = out / "communities.json"
    communities = None
    if comm_p.exists():
        try:
            communities = json.loads(comm_p.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            log(f"communities.json unreadable ({e}); public_venues not ingested")
    venues, v_rep = venue_counts(communities, geographies)

    mit_present = any(f["path"].lower().endswith((".csv", ".tab")) for f in files_of(manifest, "mit_election"))
    if mit_present:
        log("MIT election file present but no parser has been written against it yet; rep_pres_share_2024 stays not_ingested")

    def assemble(level: str, keys, cath: dict, venue_level: str | None):
        res = {}
        for k in sorted(keys):
            o = dict(acs[level][k])
            o.update(cath.get(k, {m: obs(None, None, ["not_published"]) for m in CATHOLIC_METRICS}))
            o["rep_pres_share_2024"] = obs(None, None, ["not_ingested"])
            if venues is None:
                o["public_venues"] = obs(None, None, ["not_ingested"])
            elif venue_level == "county":
                n, u = venues["county"]
                o["public_venues"] = venue_obs(n.get(k, 0), u.get(k, 0))
            elif venue_level == "state":
                n, u = venues["state"]
                o["public_venues"] = venue_obs(n.get(k, 0), u.get(k, 0))
            elif venue_level == "cbsa":
                n, u = venues["county"]
                o["public_venues"] = venue_obs(sum(n.get(c, 0) for c in members[k]), sum(u.get(c, 0) for c in members[k]))
            res[k] = {m: o[m] for m in COUNTY_METRICS}
        return res

    vintage = f"{ACS_VINTAGE_LABEL} · {USRC_VINTAGE_LABEL}"
    src_ids = ["acs5", "usrc2020", "census_geo", "communities", "mit_election"]
    county_obs = assemble("county", acs["county"].keys(), c_cath, "county")
    state_obs = assemble("state", acs["state"].keys(), s_cath, "state")
    cbsa_obs = assemble("cbsa", acs["cbsa"].keys(), b_cath, "cbsa")

    status = {"acs5": "ingested", "acs5_variables": "ingested", "census_geo": "ingested",
              "usrc2020": "ingested" if usrc else "not_ingested",
              "communities": "ingested" if venues is not None else "not_ingested"}
    sources, metrics, merge_notes = merge_registry(base_sources(manifest, status), metric_defs(resolved), registry_dir)
    checks = internal_checks(acs, usrc, geographies)
    checks.update({"usrc_county_join": {k: (v if not isinstance(v, list) else {"n": len(v), "ids": v[:40]})
                                        for k, v in c_rep.items()},
                   "usrc_cbsa_join": {"joined": b_rep["joined"], "geo_changed": len(b_rep["geo_changed"])},
                   "venues": v_rep, "counties_without_gazetteer": missing_gaz, "registry_merge": merge_notes,
                   "b01001_b12002_female_20_29_mismatches": resolved["b01001_b12002_female_20_29_mismatches"],
                   "acs_cells": resolved["labels"]})

    sizes = {}
    us = out / "us"
    sizes["us/geographies.json"] = dump(us / "geographies.json", geographies)
    for name, data in (("county", county_obs), ("state", state_obs), ("cbsa", cbsa_obs)):
        sizes[f"us/{name}_observations.json"] = dump(us / f"{name}_observations.json",
                                                     {"vintage_label": vintage, "source_ids": src_ids, "observations": data})
    sizes["us/build_checks.json"] = dump(us / "build_checks.json", {"generated_by": "pipeline/build_us.py",
                                                                     "built_at": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
                                                                     "checks": checks})
    sizes["sources.json"] = dump(out / "sources.json", {"sources": sources})
    sizes["metrics.json"] = dump(out / "metrics.json", {"metrics": metrics})
    if geo:
        shp = sorted((raw / "census_geo").glob("cb_*_us_county_20m.zip"))
        if not shp:
            raise BuildError("missing county boundary shapefile")
        gfiles = build_geo(shp[-1], keep=set(acs["county"].keys()))
        for name, g in gfiles.items():
            sizes[f"geo/{name}.json"] = dump(out / "geo" / f"{name}.json", g)
    return {"sizes": sizes, "checks": checks, "county_obs": county_obs, "geographies": geographies}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--raw", required=True, help="directory holding manifest.json and the raw source folders")
    ap.add_argument("--out", required=True, help="the venture's data/ directory")
    ap.add_argument("--registry", default=str(HERE / "registry"))
    ap.add_argument("--no-geo", action="store_true", help="skip the map paths (faster)")
    a = ap.parse_args(argv)
    try:
        res = build(Path(a.raw), Path(a.out), Path(a.registry), geo=not a.no_geo)
    except BuildError as e:
        print(f"BUILD FAILED: {e}", file=sys.stderr)
        return 2
    for k, v in sorted(res["sizes"].items()):
        print(f"  {k:40s} {v:>10,d} bytes")
    print(json.dumps({k: v for k, v in res["checks"].items() if k != "acs_cells"}, indent=1)[:4000])
    return 0


if __name__ == "__main__":
    sys.exit(main())
