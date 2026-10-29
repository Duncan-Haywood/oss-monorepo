"""SYNTHETIC test fixtures for pipeline/build_us.py.

Every number here is invented. The files mimic the *layout* of the real raw
downloads (ACS table-based Summary File, group metadata, table shells,
Gazetteer, CBSA delineation, cartographic boundary shapefile, 2020 U.S.
Religion Census workbooks) so the parser can be tested without network.
manifest.json says "synthetic": true and every place is named "Fixture ...";
build_us.py refuses to write a synthetic raw directory into data/.

    make_raw(tmp_path)                      -> a complete synthetic raw dir
    make_raw(tmp_path, usrc_county_extra=[...]) -> inject Religion Census rows
"""
from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

YEAR = 2024

# Geography: two states, four counties, one CBSA.
STATES = {"08": ("Fixture State Eight", "F8"), "09": ("Fixture State Nine", "F9")}
COUNTIES = {
    "08001": "Fixture County A", "08003": "Fixture County B", "08005": "Fixture County C",
    "09110": "Fixture Planning Region",
}
CBSA = ("99901", "Fixture Metro Area", ["08001", "08003"])

# B01001 female cells by line: (estimate, moe) per county. 08005 has a
# special-value code in the "20 years" cell.
B01001 = {
    #        026 Female      031 18-19   032 20      033 21      034 22-24    035 25-29    036 30-34
    "08001": {26: (5000, -555555555), 31: (100, 20), 32: (60, 15), 33: (55, 12), 34: (150, 30), 35: (240, 40), 36: (230, 35)},
    "08003": {26: (3000, -555555555), 31: (40, 10), 32: (0, 13), 33: (25, 9), 34: (70, 18), 35: (110, 25), 36: (100, 22)},
    "08005": {26: (800, -555555555), 31: (10, 6), 32: (-666666666, -222222222), 33: (5, 4), 34: (12, 7), 35: (20, 9), 36: (18, 8)},
    "09110": {26: (9000, -555555555), 31: (300, 40), 32: (150, 25), 33: (140, 22), 34: (400, 50), 35: (600, 60), 36: (580, 55)},
}
TOTAL_POP = {"08001": 10000, "08003": 6100, "08005": 1650, "09110": 18000}

# B12002 female 20-24 / 25-29 by marital status (never, spouse present,
# separated, spouse absent other, widowed, divorced); sums match B01001.
B12002_LINES = {  # line: (status, age)
    99: ("Never married:", "20 to 24 years"), 100: ("Never married:", "25 to 29 years"),
    101: ("Never married:", "30 to 34 years"),
    115: ("Now married:!!Married, spouse present:", "20 to 24 years"), 116: ("Now married:!!Married, spouse present:", "25 to 29 years"),
    131: ("Now married:!!Married, spouse absent:!!Separated:", "20 to 24 years"), 132: ("Now married:!!Married, spouse absent:!!Separated:", "25 to 29 years"),
    146: ("Now married:!!Married, spouse absent:!!Other:", "20 to 24 years"), 147: ("Now married:!!Married, spouse absent:!!Other:", "25 to 29 years"),
    161: ("Widowed:", "20 to 24 years"), 162: ("Widowed:", "25 to 29 years"),
    176: ("Divorced:", "20 to 24 years"), 177: ("Divorced:", "25 to 29 years"),
}


def _b12002_values(fips: str) -> dict[int, tuple[int, int]]:
    b = B01001[fips]
    w2024 = sum(b[k][0] for k in (32, 33, 34) if b[k][0] >= 0)
    w2529 = b[35][0]
    nm24, nm29 = round(w2024 * 0.8), round(w2529 * 0.5)
    rest24, rest29 = w2024 - nm24, w2529 - nm29
    vals = {99: (nm24, 20), 100: (nm29, 18), 101: (round(b[36][0] * 0.3), 15),
            115: (rest24, 10), 116: (rest29, 12)}
    for line in (131, 132, 146, 147, 161, 162, 176, 177):
        vals[line] = (0, 11)
    return vals


def _rows(level_rows: dict[str, dict[int, tuple]], table: str, lines: list[int]) -> str:
    head = ["GEO_ID"] + [f"{table}_{t}{line:03d}" for line in lines for t in ("E", "M")]
    out = ["|".join(head)]
    for gid, vals in level_rows.items():
        cells = [gid]
        for line in lines:
            e, m = vals.get(line, (0, 0))
            cells += [str(e), str(m)]
        out.append("|".join(cells))
    return "\n".join(out) + "\n"


def _sum_rows(fipss, getter, lines):
    out = {}
    for line in lines:
        e = sum(max(getter(f)[line][0], 0) for f in fipss)
        out[line] = (e, 50)
    return out


def write_acs(raw: Path, *, label_override: dict | None = None, shell_override: dict | None = None) -> None:
    d = raw / "acs5" / str(YEAR)
    sf = d / "sf"
    sf.mkdir(parents=True, exist_ok=True)
    b01001_labels = {1: "Total:", 26: "Total:!!Female:", 31: "Total:!!Female:!!18 and 19 years",
                     32: "Total:!!Female:!!20 years", 33: "Total:!!Female:!!21 years",
                     34: "Total:!!Female:!!22 to 24 years", 35: "Total:!!Female:!!25 to 29 years",
                     36: "Total:!!Female:!!30 to 34 years"}
    b12002_labels = {95: "Total:!!Female:", 96: "Total:!!Female:!!Never married:"}
    for line, (status, age) in B12002_LINES.items():
        b12002_labels[line] = f"Total:!!Female:!!{status}!!{age}"
    groups = {
        "B01001": {f"B01001_{k:03d}E": {"label": "Estimate!!" + v, "concept": "Sex by Age"} for k, v in b01001_labels.items()},
        "B12002": {f"B12002_{k:03d}E": {"label": "Estimate!!" + v, "concept": "Sex by Marital Status by Age"} for k, v in b12002_labels.items()},
        "B01003": {"B01003_001E": {"label": "Estimate!!Total", "concept": "Total Population"}},
    }
    for code, lab in (label_override or {}).items():
        groups[code.split("_")[0]][code]["label"] = lab
    for g, v in groups.items():
        (d / f"groups_{g}.json").write_text(json.dumps({"variables": v}))
    # Table shells: every label path emitted once, depth-first, with its indent.
    shell_lines = ["Table ID|Line|Indent|Unique ID|Label|Title|Universe|Type"]
    for table, labs in (("B01001", b01001_labels), ("B12002", b12002_labels), ("B01003", {1: "Total"})):
        emitted, n = set(), 0
        for line, lab in sorted(labs.items()):
            parts = lab.split("!!")
            for i in range(len(parts)):
                prefix = tuple(parts[:i + 1])
                if prefix in emitted:
                    continue
                emitted.add(prefix)
                if i == len(parts) - 1:
                    uid = f"{table}_{line:03d}"
                    text = (shell_override or {}).get(uid, parts[i])
                else:
                    n += 1
                    uid, text = f"{table}_P{n:02d}", parts[i]
                shell_lines.append(f"{table}|{line}.0|{i}|{uid}|{text}|Fixture|Fixture universe|int")
    (sf / f"ACS{YEAR}5YR_Table_Shells.txt").write_text("\n".join(shell_lines) + "\n")

    def gid_county(f):
        return f"0500000US{f}"
    b01_lines = [1, 26, 31, 32, 33, 34, 35, 36]
    county_b01 = {gid_county(f): {**v, 1: (TOTAL_POP[f], -555555555)} for f, v in B01001.items()}
    cb, _, members = CBSA
    # The CBSA row is its own estimate: deliberately NOT the sum of its counties.
    cbsa_b01 = {f"310M700US{cb}": {1: (16200, -555555555), 26: (8100, -555555555), 31: (141, 22), 32: (61, 17),
                                    33: (81, 15), 34: (222, 33), 35: (351, 45), 36: (331, 40)}}
    state_b01 = {"0400000US08": _sum_rows(["08001", "08003", "08005"], lambda f: {**B01001[f], 1: (TOTAL_POP[f], 0)}, b01_lines),
                 "0400000US09": _sum_rows(["09110"], lambda f: {**B01001[f], 1: (TOTAL_POP[f], 0)}, b01_lines)}
    us_b01 = {"0100000US": _sum_rows(list(B01001), lambda f: {**B01001[f], 1: (TOTAL_POP[f], 0)}, b01_lines)}
    allb01 = {**us_b01, **state_b01, **county_b01, **cbsa_b01}
    (sf / f"acsdt5y{YEAR}-b01001.dat").write_text(_rows(allb01, "B01001", b01_lines))
    (sf / f"acsdt5y{YEAR}-b01003.dat").write_text(_rows({g: {1: v[1]} for g, v in allb01.items()}, "B01003", [1]))
    b12_lines = sorted(B12002_LINES)
    county_b12 = {gid_county(f): _b12002_values(f) for f in B01001}
    cbsa_b12 = {f"310M700US{cb}": {**{line: (0, 11) for line in b12_lines}, 99: (300, 30), 100: (170, 25), 101: (90, 20),
                                    115: (64, 10), 116: (181, 15)}}
    st_b12 = {"0400000US08": _sum_rows(["08001", "08003", "08005"], _b12002_values, b12_lines),
              "0400000US09": _sum_rows(["09110"], _b12002_values, b12_lines)}
    us_b12 = {"0100000US": _sum_rows(list(B01001), _b12002_values, b12_lines)}
    (sf / f"acsdt5y{YEAR}-b12002.dat").write_text(_rows({**us_b12, **st_b12, **county_b12, **cbsa_b12}, "B12002", b12_lines))
    geos = ["FILEID|STUSAB|SUMLEVEL|COMPONENT|STATE|COUNTY|CBSA|GEO_ID|NAME"]
    geos.append("ACSSF|US|010|00||||0100000US|Fixture Nation")
    for st, (name, ab) in STATES.items():
        geos.append(f"ACSSF|{ab}|040|00|{st}|||0400000US{st}|{name}")
    for f, name in COUNTIES.items():
        geos.append(f"ACSSF|{STATES[f[:2]][1]}|050|00|{f[:2]}|{f[2:]}||0500000US{f}|{name}, {STATES[f[:2]][0]}")
    geos.append(f"ACSSF|US|310|00|||{cb}|310M700US{cb}|{CBSA[1]}")
    (sf / f"Geos{YEAR}5YR.txt").write_text("\n".join(geos) + "\n")


def write_census_geo(raw: Path) -> None:
    import openpyxl
    import shapefile
    g = raw / "census_geo"
    g.mkdir(parents=True, exist_ok=True)
    # Gazetteer
    lines = ["USPS\tGEOID\tANSICODE\tNAME\tALAND\tAWATER\tALAND_SQMI\tAWATER_SQMI\tINTPTLAT\tINTPTLONG"]
    for i, (f, name) in enumerate(COUNTIES.items()):
        lines.append(f"{STATES[f[:2]][1]}\t{f}\t0\t{name}\t1\t0\t{100 + i}.5\t0\t{39 + i * 0.1}\t{-105 + i * 0.3}")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(f"{YEAR}_Gaz_counties_national.txt", "\n".join(lines) + "\n")
    (g / f"{YEAR}_Gaz_counties_national.zip").write_bytes(buf.getvalue())
    # Delineation (title rows, header on row 3, footnote at the end)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "List 1"
    ws.append(["Fixture table"])
    ws.append(["List 1. FIXTURE CBSAs"])
    ws.append(["CBSA Code", "Metropolitan Division Code", "CSA Code", "CBSA Title",
               "Metropolitan/Micropolitan Statistical Area", "Metropolitan Division Title", "CSA Title",
               "County/County Equivalent", "State Name", "FIPS State Code", "FIPS County Code", "Central/Outlying County"])
    for f in CBSA[2]:
        ws.append([CBSA[0], None, None, "Fixture Metro", "Metropolitan Statistical Area", None, None,
                   COUNTIES[f], STATES[f[:2]][0], f[:2], f[2:], "Central"])
    ws.append([None])
    ws.append(["Note: fixture."])
    wb.save(g / "list1_2023.xlsx")
    # Shapefile: unit squares in lon/lat; A and B share an edge.
    squares = {"08001": (-105.0, 39.0), "08003": (-104.0, 39.0), "08005": (-104.0, 40.0), "09110": (-73.0, 41.0)}
    shp, shx, dbf = io.BytesIO(), io.BytesIO(), io.BytesIO()
    w = shapefile.Writer(shp=shp, shx=shx, dbf=dbf, shapeType=shapefile.POLYGON)
    w.field("STATEFP", "C", 2)
    w.field("COUNTYFP", "C", 3)
    w.field("GEOID", "C", 5)
    w.field("NAME", "C", 40)
    for f, (x, y) in squares.items():
        w.poly([[(x, y), (x, y + 1), (x + 1, y + 1), (x + 1, y), (x, y)]])
        w.record(f[:2], f[2:], f, COUNTIES[f])
    w.close()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("cb_2024_us_county_20m.shp", shp.getvalue())
        z.writestr("cb_2024_us_county_20m.shx", shx.getvalue())
        z.writestr("cb_2024_us_county_20m.dbf", dbf.getvalue())
    (g / "cb_2024_us_county_20m.zip").write_bytes(buf.getvalue())


USRC_COUNTY_POP = {"08001": 9800, "08003": 6000, "08005": 1600, "09001": 17500}
USRC_CATHOLIC = {  # fips: (congregations, adherents)
    "08001": (3, 1500), "08003": (1, 400), "09001": (10, 6000),
}


def write_usrc(raw: Path, *, county_extra: list | None = None) -> None:
    import openpyxl
    d = raw / "usrc2020" / "files"
    d.mkdir(parents=True, exist_ok=True)
    names = {"08001": "Fixture County A", "08003": "Fixture County B", "08005": "Fixture County C",
             "09001": "Fixture Old County"}
    sname = {"08": "Fixture State Eight", "09": "Fixture State Nine"}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2020 US Summary"
    ws.append(["2020 Population", "Congregations", "Adherents"])
    ws = wb.create_sheet("2020 State Summary")
    ws.append(["State Code", "State Name", "2020 Population", "Congregations", "Adherents"])
    ws.append(["08", sname["08"], 17400, 30, 9000])
    ws.append(["09", sname["09"], 17500, 20, 8000])
    ws = wb.create_sheet("2020 County Summary")
    ws.append(["FIPS", "State Name", "County Name", "2020 Population", "Congregations", "Adherents"])
    for f, p in USRC_COUNTY_POP.items():
        ws.append([f, sname[f[:2]], names[f], p, 10, 5000])
    ws.append(["Totals", None, None, sum(USRC_COUNTY_POP.values()), 40, 20000])
    ws = wb.create_sheet("2020 Metro Summary")
    ws.append(["CBSA Code", "Metro Name", "2020 Population", "Congregations", "Adherents"])
    ws.append([CBSA[0], CBSA[1], USRC_COUNTY_POP["08001"] + USRC_COUNTY_POP["08003"], 12, 7000])
    ws = wb.create_sheet("Copyright")
    ws.append(["Fixture copyright"])
    wb.save(d / "2020_USRC_Summaries.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2020 Group by Nation"
    ws.append(["Group Code", "Group Name", "Congregations", "Adherents"])
    ws = wb.create_sheet("2020 Group by State")
    ws.append(["StateCode", "State Name", "Group Code", "Group Name", "Congregations", "Adherents"])
    ws.append(["08", sname["08"], "081", "Catholic Church", 4, 1900])
    ws.append(["09", sname["09"], "081", "Catholic Church", 10, 6000])
    ws = wb.create_sheet("2020 Group by County")
    ws.append(["FIPS", "State Name", "County Name", "Group Code", "Group Name", "Congregations", "Adherents",
               "Adherents as % of Total Adherents", "Adherents as % of Total Population"])
    for f, (c, a) in USRC_CATHOLIC.items():
        ws.append([f, sname[f[:2]], names[f], "003", "Fixture Other Church", 2, 100, None, None])
        ws.append([f, sname[f[:2]], names[f], "081", "Catholic Church", c, a, None, None])
    for row in county_extra or []:
        ws.append(row)
    ws.append([None])
    ws.append(["Total", None, None, None, None, "=SUM(F2:F9)", "=SUM(G2:G9)", 1, 0.5])
    ws = wb.create_sheet("2020 Group by Metro")
    ws.append(["CBSA Code", "Metro Name", "Group Code", "Group Name", "Congregations", "Adherents"])
    ws.append([CBSA[0], CBSA[1], "081", "Catholic Church", 4, 1900])
    wb.save(d / "2020_USRC_Group_Detail.xlsx")


def make_raw(root: Path, *, with_usrc: bool = True, label_override: dict | None = None,
             shell_override: dict | None = None, usrc_county_extra: list | None = None) -> Path:
    raw = Path(root) / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    write_acs(raw, label_override=label_override, shell_override=shell_override)
    write_census_geo(raw)
    if with_usrc:
        write_usrc(raw, county_extra=usrc_county_extra)
    manifest = {"synthetic": True, "sources": {
        "acs5": {"status": "fetched", "files": [{"path": "acs5/2024/sf/acsdt5y2024-b01001.dat",
                                                  "url": "https://example.invalid/fixture.dat", "sha256": "0" * 64,
                                                  "retrieved_at": "2026-01-01T00:00:00+00:00"}], "errors": [], "notes": []},
        "mit_election": {"status": "failed", "files": [], "errors": ["fixture: HTTP Error 400"], "notes": []},
        "homeschool": {"status": "failed", "files": [], "errors": ["fixture: HTTP Error 403"], "notes": []},
    }}
    (raw / "manifest.json").write_text(json.dumps(manifest))
    return raw
